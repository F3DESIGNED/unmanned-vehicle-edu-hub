#!/usr/bin/env python3
"""
Export Master_Analysis_Table to CSV for R analysis
"""

import openpyxl
import csv
import glob

def main():
    # Find workbook
    wb_files = glob.glob("CTE_Research_Master_v1.0_*.xlsx")
    if not wb_files:
        print("✗ ERROR: Phase 3 workbook not found!")
        return

    wb_path = wb_files[0]
    print(f"→ Loading workbook: {wb_path}")

    wb = openpyxl.load_workbook(wb_path)

    # Export Master_Analysis_Table
    print("\n→ Exporting Master_Analysis_Table to CSV...")
    ws = wb['Master_Analysis_Table']

    with open('Master_Analysis_Table_Export.csv', 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(row)

    print(f"✓ Exported Master_Analysis_Table_Export.csv")
    print(f"  Rows: {ws.max_row}")
    print(f"  Columns: {ws.max_column}")

    # Export sample PAI sheets
    for sheet_name in ["PAI_County", "PAI_CEPD"]:
        ws = wb[sheet_name]
        filename = f"Phase3_Sample_{sheet_name}.csv"
        with open(filename, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(row)
        print(f"✓ Exported {filename}")

    print("\n✅ All exports complete")


if __name__ == "__main__":
    main()
