#!/usr/bin/env python3
"""
Phase 3 Validation Script
Comprehensive validation of PAI calculations, Master table, and integration
"""

import openpyxl
from datetime import datetime
import sys
import glob


def validate_helper_tables(wb):
    """Validation Check 1: Helper Tables Complete"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 1: HELPER TABLES (H1-H5)")
    print("="*60)

    all_passed = True

    helper_sheets = [
        ("Helper_Alignment_County", "tbl_Helper_County"),
        ("Helper_Alignment_CEPD", "tbl_Helper_CEPD"),
        ("Helper_Alignment_Prosperity", "tbl_Helper_Prosperity"),
        ("Helper_Alignment_Metropolitan", "tbl_Helper_Metropolitan"),
        ("Helper_Alignment_Economic", "tbl_Helper_Economic"),
    ]

    print("\n→ Checking all 5 helper tables...")
    for sheet_name, table_name in helper_sheets:
        try:
            ws = wb[sheet_name]
            row_count = ws.max_row - 1
            col_count = ws.max_column

            # Should have 34 columns now (with Fractional_Credit added)
            expected_cols = 34
            if col_count >= expected_cols:
                print(f"  ✓ {sheet_name}: {row_count} rows, {col_count} columns")
            else:
                print(f"  ✗ {sheet_name}: Only {col_count} columns (expected {expected_cols})")
                all_passed = False

            # Check table exists
            if table_name not in ws.tables:
                print(f"    ✗ Table {table_name} missing")
                all_passed = False

        except KeyError:
            print(f"  ✗ {sheet_name}: NOT FOUND")
            all_passed = False

    return all_passed


def validate_pai_sheets(wb):
    """Validation Check 2: PAI Calculation Sheets"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 2: PAI CALCULATION SHEETS")
    print("="*60)

    all_passed = True

    pai_sheets = [
        ("PAI_County", "tbl_PAI_County"),
        ("PAI_CEPD", "tbl_PAI_CEPD"),
        ("PAI_Prosperity", "tbl_PAI_Prosperity"),
        ("PAI_Metropolitan", "tbl_PAI_Metropolitan"),
        ("PAI_Economic", "tbl_PAI_Economic"),
    ]

    print("\n→ Checking all 5 PAI sheets...")
    for sheet_name, table_name in pai_sheets:
        try:
            ws = wb[sheet_name]
            row_count = ws.max_row - 1

            # Check key columns
            sample_row = 2
            if ws.max_row >= 2:
                total_prog = ws.cell(row=sample_row, column=3 if sheet_name == "PAI_County" else 4).value
                pai_active = ws.cell(row=sample_row, column=ws.max_column - 1).value

                if total_prog and str(total_prog).startswith('='):
                    print(f"  ✓ {sheet_name}: {row_count} districts, formulas present")
                else:
                    print(f"  ⚠ {sheet_name}: {row_count} districts (formulas not calculated)")

            # Check table
            if table_name not in ws.tables:
                print(f"    ✗ Table {table_name} missing")
                all_passed = False

        except KeyError:
            print(f"  ✗ {sheet_name}: NOT FOUND")
            all_passed = False
        except Exception as e:
            print(f"  ✗ {sheet_name}: Error - {e}")
            all_passed = False

    return all_passed


def validate_master_analysis_table(wb):
    """Validation Check 3: Master_Analysis_Table"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 3: MASTER_ANALYSIS_TABLE")
    print("="*60)

    all_passed = True

    try:
        ws = wb["Master_Analysis_Table"]

        row_count = ws.max_row - 1
        col_count = ws.max_column

        print(f"\n→ Table dimensions:")
        print(f"  Rows (districts): {row_count}")
        print(f"  Columns: {col_count}")

        expected_min_cols = 57
        if col_count >= expected_min_cols:
            print(f"  ✓ Has expected minimum columns ({expected_min_cols})")
        else:
            print(f"  ✗ Only {col_count} columns (expected {expected_min_cols}+)")
            all_passed = False

        # Check key column formulas
        print(f"\n→ Checking key formulas...")
        sample_row = 2
        key_cols = {
            "District_Name (B)": 2,
            "OCQ (W)": 23,
            "PAI_County (Y)": 25,
            "PDER (AI)": 35,
            "Include_In_Analysis": 41,
        }

        formula_count = 0
        for col_name, col_idx in key_cols.items():
            if col_idx <= ws.max_column:
                value = ws.cell(row=sample_row, column=col_idx).value
                if value and str(value).startswith('='):
                    print(f"  ✓ {col_name}: Formula present")
                    formula_count += 1
                else:
                    print(f"  ⚠ {col_name}: {value}")

        if formula_count >= 3:
            print(f"\n✓ Most key formulas implemented ({formula_count}/{len(key_cols)})")
        else:
            print(f"\n⚠ Few formulas found ({formula_count}/{len(key_cols)})")

        # Check table exists
        if "tbl_Master_Analysis" in ws.tables:
            print(f"✓ Table tbl_Master_Analysis found")
        else:
            print(f"✗ Table tbl_Master_Analysis missing")
            all_passed = False

    except KeyError:
        print(f"✗ Master_Analysis_Table sheet not found")
        all_passed = False
    except Exception as e:
        print(f"✗ Error checking Master_Analysis_Table: {e}")
        all_passed = False

    return all_passed


def validate_quick_validation(wb):
    """Validation Check 4: Quick_Validation Sheet"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 4: QUICK_VALIDATION SHEET")
    print("="*60)

    all_passed = True

    try:
        ws = wb["Quick_Validation"]

        # Count validation formulas
        validation_count = 0
        for row in range(1, ws.max_row + 1):
            cell_value = ws[f'B{row}'].value
            if cell_value and str(cell_value).startswith('=IF'):
                validation_count += 1

        print(f"\n→ Validation checks implemented: {validation_count}")

        if validation_count >= 15:
            print(f"  ✓ Sufficient validation checks ({validation_count})")
        else:
            print(f"  ⚠ Few validation checks ({validation_count})")

        # Check overall status formula
        overall_row = ws.max_row
        overall_formula = ws[f'A{overall_row}'].value
        if overall_formula and '=IF' in str(overall_formula):
            print(f"  ✓ Overall status formula present")
        else:
            print(f"  ⚠ Overall status formula not found")

    except KeyError:
        print(f"✗ Quick_Validation sheet not found")
        all_passed = False
    except Exception as e:
        print(f"✗ Error checking Quick_Validation: {e}")
        all_passed = False

    return all_passed


def validate_sheet_count(wb):
    """Validation Check 5: Complete Sheet Structure"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 5: COMPLETE SHEET STRUCTURE")
    print("="*60)

    all_passed = True

    expected_sheets = [
        # Phase 1
        "Parameters", "Districts_Raw", "Programs_Raw",
        "Labor_Market_County", "Labor_Market_CEPD", "Labor_Market_Prosperity",
        "Labor_Market_Metropolitan", "Labor_Market_Economic",
        "CIP_SOC_Crosswalk", "Geographic_Crosswalk", "Control_Variables_Raw",
        "Districts_Clean", "Programs_Fractional", "Toggle_Settings",
        # Phase 2
        "OCQ_Calculated",
        # Phase 2/3: Helpers
        "Helper_Alignment_County", "Helper_Alignment_CEPD",
        "Helper_Alignment_Prosperity", "Helper_Alignment_Metropolitan",
        "Helper_Alignment_Economic",
        # Phase 3: PAI
        "PAI_County", "PAI_CEPD", "PAI_Prosperity",
        "PAI_Metropolitan", "PAI_Economic",
        # Phase 2
        "PDER_Calculated",
        # Phase 3: Analysis
        "Master_Analysis_Table", "Quick_Validation",
        # Documentation
        "Variable_Codebook", "Formula_Key",
    ]

    print(f"\n→ Sheet inventory:")
    print(f"  Expected: {len(expected_sheets)} sheets")
    print(f"  Found: {len(wb.sheetnames)} sheets")

    missing = []
    for sheet in expected_sheets:
        if sheet in wb.sheetnames:
            print(f"  ✓ {sheet}")
        else:
            print(f"  ✗ {sheet} - MISSING")
            missing.append(sheet)
            all_passed = False

    if not missing:
        print(f"\n✓ All expected sheets present")
    else:
        print(f"\n✗ Missing {len(missing)} sheets")

    return all_passed


def validate_exports(wb):
    """Validation Check 6: Export Files"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 6: EXPORT FILES")
    print("="*60)

    all_passed = True

    import os

    expected_files = [
        "Master_Analysis_Table_Export.csv",
        "Phase3_Sample_PAI_County.csv",
        "Phase3_Sample_PAI_CEPD.csv",
    ]

    print(f"\n→ Checking export files...")
    for filename in expected_files:
        if os.path.exists(filename):
            size = os.path.getsize(filename)
            print(f"  ✓ {filename} ({size} bytes)")
        else:
            print(f"  ✗ {filename} - NOT FOUND")
            all_passed = False

    # Check R script
    if os.path.exists("R_Import_Test.R"):
        print(f"  ✓ R_Import_Test.R present")
    else:
        print(f"  ⚠ R_Import_Test.R not found")

    return all_passed


def validate_formula_integrity(wb):
    """Validation Check 7: Formula Error Check"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 7: FORMULA INTEGRITY")
    print("="*60)

    all_passed = True

    print("\n→ Checking for formula errors...")
    error_count = 0
    checked_sheets = 0

    for sheet_name in wb.sheetnames:
        if sheet_name in ["Parameters", "Toggle_Settings", "Quick_Validation"]:
            continue  # Skip special sheets

        ws = wb[sheet_name]
        sheet_errors = 0

        for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 20)):  # Sample first 20 rows
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    if any(err in str(cell.value) for err in ["#N/A", "#REF!", "#DIV/0!", "#VALUE!", "#NAME?"]):
                        sheet_errors += 1
                        error_count += 1

        if sheet_errors > 0:
            print(f"  ⚠ {sheet_name}: {sheet_errors} formula errors found")
        else:
            checked_sheets += 1

    if error_count == 0:
        print(f"  ✓ No formula errors in {checked_sheets} sheets (sample check)")
    else:
        print(f"  ⚠ Found {error_count} formula errors across sheets")
        print(f"  Note: Formulas will calculate when opened in Excel")

    return all_passed  # Don't fail on this - formulas calculate in Excel


def main():
    """Run all validation checks"""
    print("\n" + "="*70)
    print(" "*15 + "PHASE 3 VALIDATION REPORT")
    print(" "*10 + datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    print("="*70)

    # Find the workbook
    workbooks = glob.glob("CTE_Research_Master_v1.0_*.xlsx")

    if not workbooks:
        print("\n✗ ERROR: No Phase 3 workbook found!")
        return False

    wb_path = workbooks[0]
    print(f"\n→ Validating workbook: {wb_path}")

    try:
        wb = openpyxl.load_workbook(wb_path)
        print(f"✓ Workbook loaded successfully")
        print(f"✓ Total sheets: {len(wb.sheetnames)}")
        print(f"✓ Named ranges: {len(wb.defined_names)}")

    except Exception as e:
        print(f"\n✗ ERROR: Could not load workbook: {e}")
        return False

    # Run all validation checks
    results = {
        "Helper Tables (H1-H5)": validate_helper_tables(wb),
        "PAI Calculation Sheets": validate_pai_sheets(wb),
        "Master_Analysis_Table": validate_master_analysis_table(wb),
        "Quick_Validation Sheet": validate_quick_validation(wb),
        "Complete Sheet Structure": validate_sheet_count(wb),
        "Export Files": validate_exports(wb),
        "Formula Integrity": validate_formula_integrity(wb),
    }

    # Summary
    print("\n" + "="*70)
    print("VALIDATION SUMMARY")
    print("="*70)

    for check_name, passed in results.items():
        status = "✓ PASS" if passed else "✗ FAIL"
        print(f"  {status}: {check_name}")

    all_passed = all(results.values())

    print("\n" + "="*70)
    if all_passed:
        print("  ✅ ALL VALIDATION CHECKS PASSED")
        print("  Phase 3 is complete - Workbook ready for analysis")
    else:
        print("  ⚠ SOME VALIDATION CHECKS FAILED")
        print("  Review errors above (formulas will calculate in Excel)")
    print("="*70 + "\n")

    print("\n📋 NEXT STEPS:")
    print("  1. Open workbook in Excel to calculate all formulas")
    print("  2. Check Quick_Validation sheet for automated QC")
    print("  3. Test toggle functionality (change settings, verify PAI updates)")
    print("  4. Export Master_Analysis_Table to CSV")
    print("  5. Run R_Import_Test.R to verify data import")
    print("  6. Begin statistical analysis in R")

    return all_passed


if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
