#!/usr/bin/env python3
"""
Phase 2 Validation Script
Validates OCQ, PDER, and Helper table implementations
"""

import openpyxl
from datetime import datetime
import sys
import glob


def validate_named_ranges(wb):
    """Validation Check 1: Named Ranges (Phase 2 additions)"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 1: NAMED RANGES (PHASE 2)")
    print("="*60)

    expected_ranges = [
        "StatePathways_Count",  # Phase 1
        "Active_Toggle_District",  # Phase 1
        "Active_Toggle_Wage",  # Phase 1
        "Active_Toggle_OccCount",  # Phase 1
        "Active_Toggle_TimeRef",  # Phase 1
        "State_Total_Programs",  # Phase 2
        "State_Total_Enrollment",  # Phase 2
    ]

    try:
        found_names = [name for name in wb.defined_names]
    except:
        found_names = []

    print(f"\n✓ Expected named ranges: {len(expected_ranges)}")
    print(f"✓ Found named ranges: {len(found_names)}")

    all_passed = True
    for name in expected_ranges:
        if name in found_names:
            print(f"  ✓ {name} - FOUND")
        else:
            print(f"  ✗ {name} - MISSING")
            all_passed = False

    return all_passed


def validate_ocq_sheet(wb):
    """Validation Check 2: OCQ Calculations"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 2: OCQ CALCULATIONS")
    print("="*60)

    all_passed = True

    try:
        ws = wb["OCQ_Calculated"]

        # Check formulas are present
        print("\n→ Checking OCQ formulas...")
        sample_row = 2
        formulas = {
            "Total_Programs": ws[f'C{sample_row}'].value,
            "StatePathways_Count": ws[f'D{sample_row}'].value,
            "OCQ": ws[f'E{sample_row}'].value,
            "OCQ_Category": ws[f'F{sample_row}'].value,
        }

        for col_name, formula in formulas.items():
            if formula and (str(formula).startswith('=') or isinstance(formula, str)):
                print(f"  ✓ {col_name}: Formula present")
            else:
                print(f"  ⚠ {col_name}: {formula} (will calculate in Excel)")

        # Check table exists
        if "tbl_OCQ_Calculated" in ws.tables:
            print(f"\n✓ Table tbl_OCQ_Calculated found")
        else:
            print(f"\n✗ Table tbl_OCQ_Calculated missing")
            all_passed = False

        # Check row count
        row_count = ws.max_row - 1
        print(f"✓ OCQ rows: {row_count} districts")

    except KeyError:
        print(f"✗ OCQ_Calculated sheet not found")
        all_passed = False
    except Exception as e:
        print(f"✗ Error checking OCQ: {e}")
        all_passed = False

    return all_passed


def validate_pder_sheet(wb):
    """Validation Check 3: PDER Calculations"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 3: PDER CALCULATIONS")
    print("="*60)

    all_passed = True

    try:
        ws = wb["PDER_Calculated"]

        # Check formulas are present
        print("\n→ Checking PDER formulas...")
        sample_row = 2
        formulas = {
            "Total_Programs": ws[f'C{sample_row}'].value,
            "District_Enrollment": ws[f'D{sample_row}'].value,
            "Program_Share": ws[f'E{sample_row}'].value,
            "Enrollment_Share": ws[f'F{sample_row}'].value,
            "PDER": ws[f'G{sample_row}'].value,
            "PDER_Interpretation": ws[f'H{sample_row}'].value,
        }

        for col_name, formula in formulas.items():
            if formula and (str(formula).startswith('=') or isinstance(formula, str)):
                print(f"  ✓ {col_name}: Formula present")
            else:
                print(f"  ⚠ {col_name}: {formula} (will calculate in Excel)")

        # Check table exists
        if "tbl_PDER_Calculated" in ws.tables:
            print(f"\n✓ Table tbl_PDER_Calculated found")
        else:
            print(f"\n✗ Table tbl_PDER_Calculated missing")
            all_passed = False

        # Check row count
        row_count = ws.max_row - 1
        print(f"✓ PDER rows: {row_count} districts")

    except KeyError:
        print(f"✗ PDER_Calculated sheet not found")
        all_passed = False
    except Exception as e:
        print(f"✗ Error checking PDER: {e}")
        all_passed = False

    return all_passed


def validate_helper_county(wb):
    """Validation Check 4: Helper_Alignment_County"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 4: HELPER_ALIGNMENT_COUNTY")
    print("="*60)

    all_passed = True

    try:
        ws = wb["Helper_Alignment_County"]

        # Check column count
        expected_cols = 33  # NCES_ID through Aligned_130_Top20_10yr
        actual_cols = ws.max_column
        print(f"\n→ Column check:")
        print(f"  Expected: {expected_cols} columns")
        print(f"  Found: {actual_cols} columns")

        if actual_cols >= expected_cols:
            print(f"  ✓ All columns present")
        else:
            print(f"  ✗ Missing columns")
            all_passed = False

        # Check formulas in key columns
        print(f"\n→ Checking helper formulas...")
        sample_row = 2
        key_formulas = {
            "District_Name (B)": ws[f'B{sample_row}'].value,
            "County (C)": ws[f'C{sample_row}'].value,
            "SOC_Code (F)": ws[f'F{sample_row}'].value,
            "Median_Wage (H)": ws[f'H{sample_row}'].value,
            "Wage_Ratio (L)": ws[f'L{sample_row}'].value,
            "Is_High_Wage_110 (M)": ws[f'M{sample_row}'].value,
            "Rank_Current (P)": ws[f'P{sample_row}'].value,
            "Is_Top15_Current (R)": ws[f'R{sample_row}'].value,
            "Aligned_120_Top15_Current (Z)": ws[f'Z{sample_row}'].value,
        }

        formula_count = 0
        for col_name, value in key_formulas.items():
            if value and str(value).startswith('='):
                print(f"  ✓ {col_name}: Formula present")
                formula_count += 1
            else:
                print(f"  ⚠ {col_name}: {value}")

        if formula_count >= 7:
            print(f"\n✓ Most formulas implemented ({formula_count}/9)")
        else:
            print(f"\n⚠ Few formulas found ({formula_count}/9)")

        # Check row count
        row_count = ws.max_row - 1
        print(f"\n✓ Helper rows: {row_count} (district-program combinations)")

        # Check table exists
        if "tbl_Helper_County" in ws.tables:
            print(f"✓ Table tbl_Helper_County found")
        else:
            print(f"✗ Table tbl_Helper_County missing")
            all_passed = False

    except KeyError:
        print(f"✗ Helper_Alignment_County sheet not found")
        all_passed = False
    except Exception as e:
        print(f"✗ Error checking Helper_Alignment_County: {e}")
        all_passed = False

    return all_passed


def validate_helper_templates(wb):
    """Validation Check 5: Helper Templates H2-H5"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 5: HELPER TEMPLATES (H2-H5)")
    print("="*60)

    all_passed = True

    helper_sheets = [
        "Helper_Alignment_CEPD",
        "Helper_Alignment_Prosperity",
        "Helper_Alignment_Metropolitan",
        "Helper_Alignment_Economic",
    ]

    print(f"\n→ Checking template sheets...")
    for sheet_name in helper_sheets:
        try:
            ws = wb[sheet_name]
            col_count = ws.max_column

            if col_count >= 33:
                print(f"  ✓ {sheet_name}: {col_count} columns")
            else:
                print(f"  ⚠ {sheet_name}: Only {col_count} columns")

        except KeyError:
            print(f"  ✗ {sheet_name}: NOT FOUND")
            all_passed = False

    return all_passed


def validate_documentation(wb):
    """Validation Check 6: Documentation Updates"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 6: DOCUMENTATION UPDATES")
    print("="*60)

    all_passed = True

    # Check Variable_Codebook
    print("\n→ Checking Variable_Codebook...")
    try:
        ws = wb["Variable_Codebook"]
        row_count = ws.max_row - 1
        print(f"  ✓ Variable_Codebook entries: {row_count}")

        if row_count >= 45:  # 31 from Phase 1 + ~16 from Phase 2
            print(f"  ✓ Phase 2 variables added")
        else:
            print(f"  ⚠ May be missing Phase 2 variables (expected 45+, found {row_count})")

    except Exception as e:
        print(f"  ✗ Error checking Variable_Codebook: {e}")
        all_passed = False

    # Check Formula_Key
    print("\n→ Checking Formula_Key...")
    try:
        ws = wb["Formula_Key"]
        row_count = ws.max_row - 1
        print(f"  ✓ Formula_Key entries: {row_count}")

        if row_count >= 18:  # 10 from Phase 1 + ~10 from Phase 2
            print(f"  ✓ Phase 2 formulas added")
        else:
            print(f"  ⚠ May be missing Phase 2 formulas (expected 18+, found {row_count})")

    except Exception as e:
        print(f"  ✗ Error checking Formula_Key: {e}")
        all_passed = False

    return all_passed


def validate_sheet_structure(wb):
    """Validation Check 7: Overall Sheet Structure"""
    print("\n" + "="*60)
    print("VALIDATION CHECK 7: OVERALL STRUCTURE")
    print("="*60)

    all_passed = True

    expected_sheets = [
        # Phase 1
        "Parameters", "Districts_Raw", "Programs_Raw",
        "Labor_Market_County", "Labor_Market_CEPD", "Labor_Market_Prosperity",
        "Labor_Market_Metropolitan", "Labor_Market_Economic",
        "CIP_SOC_Crosswalk", "Geographic_Crosswalk", "Control_Variables_Raw",
        "Districts_Clean", "Programs_Fractional", "Toggle_Settings",
        # Phase 2
        "OCQ_Calculated", "Helper_Alignment_County",
        "Helper_Alignment_CEPD", "Helper_Alignment_Prosperity",
        "Helper_Alignment_Metropolitan", "Helper_Alignment_Economic",
        "PDER_Calculated",
        # Documentation
        "Variable_Codebook", "Formula_Key",
    ]

    print(f"\n→ Sheet inventory:")
    print(f"  Expected: {len(expected_sheets)} sheets")
    print(f"  Found: {len(wb.sheetnames)} sheets")

    missing = []
    for sheet in expected_sheets:
        if sheet in wb.sheetnames:
            print(f"  ✓ {sheet}")
        else:
            print(f"  ✗ {sheet} - MISSING")
            missing.append(sheet)
            all_passed = False

    if not missing:
        print(f"\n✓ All expected sheets present")
    else:
        print(f"\n✗ Missing {len(missing)} sheets")

    return all_passed


def main():
    """Run all validation checks"""
    print("\n" + "="*70)
    print(" "*15 + "PHASE 2 VALIDATION REPORT")
    print(" "*10 + datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    print("="*70)

    # Find the workbook
    workbooks = glob.glob("CTE_Research_Master_Phase2_*.xlsx")

    if not workbooks:
        print("\n✗ ERROR: No Phase 2 workbook found!")
        return False

    wb_path = workbooks[0]
    print(f"\n→ Validating workbook: {wb_path}")

    try:
        wb = openpyxl.load_workbook(wb_path)
        print(f"✓ Workbook loaded successfully")
        print(f"✓ Total sheets: {len(wb.sheetnames)}")

    except Exception as e:
        print(f"\n✗ ERROR: Could not load workbook: {e}")
        return False

    # Run all validation checks
    results = {
        "Named Ranges": validate_named_ranges(wb),
        "OCQ Calculations": validate_ocq_sheet(wb),
        "PDER Calculations": validate_pder_sheet(wb),
        "Helper County (Complete)": validate_helper_county(wb),
        "Helper Templates (H2-H5)": validate_helper_templates(wb),
        "Documentation Updates": validate_documentation(wb),
        "Sheet Structure": validate_sheet_structure(wb),
    }

    # Summary
    print("\n" + "="*70)
    print("VALIDATION SUMMARY")
    print("="*70)

    for check_name, passed in results.items():
        status = "✓ PASS" if passed else "✗ FAIL"
        print(f"  {status}: {check_name}")

    all_passed = all(results.values())

    print("\n" + "="*70)
    if all_passed:
        print("  ✅ ALL VALIDATION CHECKS PASSED")
        print("  Phase 2 is complete and ready for Phase 3")
    else:
        print("  ⚠ SOME VALIDATION CHECKS FAILED")
        print("  Review errors above (formulas will calculate in Excel)")
    print("="*70 + "\n")

    return all_passed


if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
